"""按文件扩展名提取文本。图片/视频不在本模块处理（交 VL）。"""
from pathlib import Path

TEXT_EXTS = {".txt", ".md", ".markdown"}
DOCX_EXTS = {".docx"}
PDF_EXTS = {".pdf"}
XLSX_EXTS = {".xlsx", ".xlsm"}
PPTX_EXTS = {".pptx"}
IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".webp", ".bmp"}
VIDEO_EXTS = {".mp4", ".mov", ".webm", ".mkv", ".avi"}

ALL_TEXT_EXTS = TEXT_EXTS | DOCX_EXTS | PDF_EXTS | XLSX_EXTS | PPTX_EXTS


def classify(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in ALL_TEXT_EXTS:
        return "text"
    if ext in IMAGE_EXTS:
        return "image"
    if ext in VIDEO_EXTS:
        return "video"
    return "other"


def extract(path: Path) -> str:
    """返回单个文件里提取到的纯文本（图片/视频返回空串）。"""
    ext = path.suffix.lower()
    try:
        if ext in TEXT_EXTS:
            return path.read_text(encoding="utf-8", errors="ignore")
        if ext in DOCX_EXTS:
            return _extract_docx(path)
        if ext in PDF_EXTS:
            return _extract_pdf(path)
        if ext in XLSX_EXTS:
            return _extract_xlsx(path)
        if ext in PPTX_EXTS:
            return _extract_pptx(path)
    except Exception as e:  # 提取失败不阻断流水线
        return f"[提取失败: {type(e).__name__}: {e}]"
    return ""


def _extract_docx(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


def _extract_pdf(path: Path) -> str:
    import pdfplumber
    out = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            out.append(page.extract_text() or "")
    return "\n".join(out)


def _extract_xlsx(path: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(str(path), data_only=True, read_only=True)
    parts = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None and str(c).strip()]
            if cells:
                parts.append("\t".join(cells))
    return "\n".join(parts)


def _extract_pptx(path: Path) -> str:
    from pptx import Presentation
    prs = Presentation(str(path))
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"--- slide {i} ---")
        for shape in slide.shapes:
            if shape.has_text_frame:
                txt = shape.text_frame.text
                if txt.strip():
                    parts.append(txt)
    return "\n".join(parts)


def walk_input(input_path: str):
    """遍历输入路径，返回 {text:[(path,type,text)], image:[path], video:[path], other:[path]}。"""
    p = Path(input_path).expanduser().resolve()
    if p.is_file():
        files = [p]
    else:
        files = [f for f in p.rglob("*") if f.is_file()]
    out = {"text": [], "image": [], "video": [], "other": []}
    for f in files:
        kind = classify(f)
        if kind == "text":
            out["text"].append((f, kind, extract(f)))
        elif kind == "image":
            out["image"].append(f)
        elif kind == "video":
            out["video"].append(f)
        else:
            out["other"].append(f)
    return out
